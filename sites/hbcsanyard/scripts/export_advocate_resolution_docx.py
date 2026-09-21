#!/usr/bin/env python3
"""Write the certified EC resolution (advocate appointment + 50/50 fee split) as Word."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x2A, 0x56)
GREEN = RGBColor(0x1A, 0x6B, 0x3A)
INK = RGBColor(0x12, 0x23, 0x3F)
MUTED = RGBColor(0x5A, 0x6A, 0x80)

SITE = Path(__file__).resolve().parent.parent
DEFAULT_OUT = SITE / "documents" / "ResolutionAppointmentOfNewAdvocate.docx"
SEAL = SITE / "assets" / "mhws-logo" / "mhws-logo-seal-cert.png"


def _set_run(run, *, size=11, bold=False, italic=False, color=INK, name="Times New Roman"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _p(doc, text="", *, size=11, bold=False, italic=False, color=INK, align="justify", space_after=8, space_before=0, first_line=None):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if text:
        run = para.add_run(text)
        _set_run(run, size=size, bold=bold, italic=italic, color=color)
    return para


def _add_runs(para, parts):
    """parts: list of (text, kwargs for _set_run)."""
    for text, kwargs in parts:
        run = para.add_run(text)
        _set_run(run, **kwargs)
    return para


def _bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "0B2A56")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.8)

    if SEAL.exists():
        head = doc.add_paragraph()
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        head.paragraph_format.space_after = Pt(2)
        run = head.add_run()
        run.add_picture(str(SEAL), width=Cm(2.4))

    t = _p(doc, "MANDI HOUSING WELFARE SOCIETY", size=16, bold=True, color=NAVY, align="center", space_after=2)
    _bottom_border(t)
    _p(doc, "Himuda Housing Colony Sanyard", size=11, bold=True, color=GREEN, align="center", space_after=1, space_before=6)
    _p(
        doc,
        "Housing Colony Sanyard, Mandi (H.P.) 175001  ·  Registration No. 467 dated 21/07/2012",
        size=9,
        color=MUTED,
        align="center",
        space_after=1,
    )
    _p(
        doc,
        "housingcolonysanyard@gmail.com  ·  housingcolonysanyard.in",
        size=9,
        color=MUTED,
        align="center",
        space_after=14,
    )

    _p(doc, "CERTIFIED TRUE COPY OF RESOLUTION", size=13, bold=True, color=NAVY, align="center", space_after=10)

    intro = _p(doc, "", space_after=8)
    _add_runs(
        intro,
        [
            ("Passed by the ", {"size": 11}),
            ("Executive Committee", {"size": 11, "bold": True}),
            (" of Mandi Housing Welfare Society at its meeting held on ", {"size": 11}),
            ("15/08/2026", {"size": 11, "bold": True}),
            (" at Housing Colony Sanyard, Mandi (H.P.), with the sharing of professional fee as recorded in the minutes of the Executive Committee meeting held on ", {"size": 11}),
            ("12/09/2026", {"size": 11, "bold": True}),
            (" (Register No. ", {"size": 11}),
            ("1/2026", {"size": 11, "bold": True}),
            (").", {"size": 11}),
        ],
    )

    subj = _p(doc, "", space_after=8)
    _add_runs(
        subj,
        [
            ("Subject: ", {"size": 11, "bold": True}),
            (
                "Appointment of Advocate Mr. Shailesh Sharma for conduct of pending civil litigation; "
                "engagement of Mr. Bimal Sharma for continued guidance; sharing of professional fee "
                "between the Society and Shri B.C. Sharma.",
                {"size": 11},
            ),
        ],
    )

    rec = _p(doc, "", space_after=8)
    _add_runs(
        rec,
        [
            (
                "The Committee considered that the civil suit concerning the ",
                {"size": 11},
            ),
            ("right of path / link road through Himuda Housing Colony Sanyard", {"size": 11, "italic": True}),
            (
                " (pending before the Senior Civil Judge, Mandi, Civil Suit No. ≈ 086/2023, or as numbered on the court record) "
                "is of vital importance to the Society and its members; that the case requires ",
                {"size": 11},
            ),
            ("active legal representation", {"size": 11, "italic": True}),
            (
                " for further hearings, filings, and proceedings; that ",
                {"size": 11},
            ),
            ("Mr. Bimal Sharma, Advocate", {"size": 11, "bold": True}),
            (", has hitherto represented the Society in the matter; and that ", {"size": 11}),
            ("Mr. Shailesh Sharma, Advocate", {"size": 11, "bold": True}),
            (", has been ", {"size": 11}),
            ("confirmed", {"size": 11, "italic": True}),
            (" by the Committee to take forward the conduct of the case.", {"size": 11}),
        ],
    )

    r1 = _p(doc, "", space_after=8)
    _add_runs(
        r1,
        [
            ("RESOLVED THAT ", {"size": 11, "bold": True}),
            ("Mr. Shailesh Sharma, Advocate, be and is hereby ", {"size": 11}),
            ("appointed and confirmed", {"size": 11, "italic": True}),
            (
                " as the counsel of the Mandi Housing Welfare Society to ",
                {"size": 11},
            ),
            ("take over, conduct, and pursue", {"size": 11, "italic": True}),
            (
                " the aforesaid pending civil suit and all connected applications, hearings, and proceedings "
                "before the Hon’ble Court and any other forum as may be required in connection therewith.",
                {"size": 11},
            ),
        ],
    )

    r2 = _p(doc, "", space_after=8)
    _add_runs(
        r2,
        [
            ("RESOLVED FURTHER THAT ", {"size": 11, "bold": True}),
            (
                "for the conduct of the said case, Mr. Shailesh Sharma, Advocate, shall be paid a ",
                {"size": 11},
            ),
            ("lump-sum professional fee of Rs. 50,000/- (Rupees Fifty Thousand only)", {"size": 11, "bold": True}),
            (" for this case.", {"size": 11}),
        ],
    )

    r3 = _p(doc, "", space_after=8)
    _add_runs(
        r3,
        [
            ("RESOLVED FURTHER THAT ", {"size": 11, "bold": True}),
            (
                "as decided and accepted at the Executive Committee meeting held on 12/09/2026 "
                "(Register No. 1/2026; resolution passed For: 4, Against: 0, Abstain: 0), the said "
                "professional fee of Rs. 50,000/- shall be borne equally: ",
                {"size": 11},
            ),
            ("50% (Rs. 25,000/-) by the Society", {"size": 11, "bold": True}),
            (
                " from the Society’s funds (Bank of Baroda, Mandi Branch, Account No. 09640100004511, "
                "IFSC BARB0MANDIX), in accordance with the Society’s rules and banking instructions; and ",
                {"size": 11},
            ),
            ("50% (Rs. 25,000/-) by Shri B.C. Sharma", {"size": 11, "bold": True}),
            (
                ", who is also a party in the case. Shri B.C. Sharma has accepted this proposal. "
                "It is recorded that Rs. 1,500/- has already been paid to the counsel by Shri B.C. Sharma "
                "toward his share; the pending amount shall be paid in due course. Court fees, stamps, process, "
                "certified copies and similar out-of-pocket expenses of the case, if any, shall be borne "
                "separately against bills and are not included in the said professional fee of Rs. 50,000/-.",
                {"size": 11},
            ),
        ],
    )

    r4 = _p(doc, "", space_after=8)
    _add_runs(
        r4,
        [
            ("RESOLVED FURTHER THAT ", {"size": 11, "bold": True}),
            ("the conduct of the case shall ", {"size": 11}),
            ("transfer from Mr. Bimal Sharma, Advocate, to Mr. Shailesh Sharma, Advocate", {"size": 11, "italic": True}),
            (
                ", for all future court work and formal representation; and that Mr. Bimal Sharma, Advocate, "
                "having agreed thereto, shall ",
                {"size": 11},
            ),
            ("continue to provide guidance and advice", {"size": 11, "italic": True}),
            (
                " to the Society in the same matter as and when required, on terms to be mutually agreed "
                "and recorded separately if necessary.",
                {"size": 11},
            ),
        ],
    )

    r5 = _p(doc, "", space_after=10)
    _add_runs(
        r5,
        [
            ("RESOLVED FURTHER THAT ", {"size": 11, "bold": True}),
            ("the President and the General Secretary of the Society be and are hereby jointly and severally ", {"size": 11}),
            ("authorised", {"size": 11, "italic": True}),
            (
                " to: sign the vakalatnama / power to act, engagement letter, and all pleadings, affidavits, "
                "applications, and correspondence; brief and instruct Mr. Shailesh Sharma, Advocate, and "
                "Mr. Bimal Sharma, Advocate, as appropriate; release payment of the Society’s share of "
                "Rs. 25,000/- and any incidental lawful expenses relating to the case; do all other acts "
                "necessary to give effect to this resolution.",
                {"size": 11},
            ),
        ],
    )

    cert = _p(doc, "", space_after=10)
    _add_runs(
        cert,
        [
            ("Certified that the above is a ", {"size": 11}),
            ("true extract", {"size": 11, "italic": True}),
            (
                " of the resolution passed by the Executive Committee, that it is in force, and that it has "
                "not been modified or rescinded.",
                {"size": 11},
            ),
        ],
    )

    place = _p(doc, "", align="left", space_after=4)
    _add_runs(place, [("Place: Mandi (H.P.)", {"size": 11})])
    date = _p(doc, "", align="left", space_after=14)
    _add_runs(date, [("Date: ______________", {"size": 11})])

    _p(doc, "For Mandi Housing Welfare Society", size=11, bold=True, align="left", space_after=28)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    for cell, name, role in (
        (left, "Anup Vaidya", "President"),
        (right, "Vijay Kumar Sharma", "General Secretary"),
    ):
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p1.add_run("______________________________")
        _set_run(r, size=11)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(name)
        _set_run(r, size=11, bold=True, color=NAVY)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p3.add_run(role)
        _set_run(r, size=10, color=MUTED)
        p4 = cell.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p4.add_run("(Society seal)")
        _set_run(r, size=9, italic=True, color=MUTED)

    note = _p(doc, "", align="left", space_before=16, space_after=0)
    _add_runs(
        note,
        [
            (
                "Note: Fee-sharing as per EC MOM dated 12/09/2026 (Register 1/2026). "
                "Appointment of counsel as per EC meeting dated 15/08/2026.",
                {"size": 8, "italic": True, "color": MUTED},
            ),
        ],
    )
    return doc


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(str(args.output))
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
